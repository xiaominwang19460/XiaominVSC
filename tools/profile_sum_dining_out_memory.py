"""Create a large XLSX and profile memory used by openpyxl when loading it.

Run from the repository root. Installs: openpyxl (from requirements.txt).
"""
import tracemalloc
import time
from pathlib import Path
from openpyxl import Workbook, load_workbook


def make_workbook(path: Path, sheets: int = 5, rows_per_sheet: int = 2000) -> None:
    wb = Workbook()
    # remove default sheet
    default = wb.active
    wb.remove(default)
    for i in range(sheets):
        ws = wb.create_sheet(title=f"{i+1:02d} 2024")
        # create 20 header rows to match consumer code
        for _ in range(20):
            ws.append([None, None, None, None])
        for r in range(rows_per_sheet):
            # column C has numeric values, column D sometimes matches target
            value_c = float(r % 100) + 0.5
            value_d = "Dinning Out" if (r % 10 == 0) else "Other"
            ws.append([None, None, value_c, value_d])
    wb.save(path)


def measure_load(path: Path, read_only: bool) -> None:
    print(f"\nMeasuring load_workbook(read_only={read_only})")
    tracemalloc.start()
    t0 = time.time()
    wb = load_workbook(path, data_only=True, read_only=read_only)
    t1 = time.time()
    current, peak = tracemalloc.get_traced_memory()
    print(f"Loaded workbook in {t1 - t0:.2f}s — tracemalloc current={current/1024**2:.2f} MiB peak={peak/1024**2:.2f} MiB")
    # Access one sheet and iterate a few rows to ensure consumption
    sheet = next((s for s in wb.worksheets if s.title and len(s.title) > 0), None)
    if sheet is not None:
        it = sheet.iter_rows(min_row=21, max_row=30, values_only=True)
        for _ in it:
            pass
    try:
        wb.close()
    except Exception:
        pass
    tracemalloc.stop()


def main() -> None:
    repo_root = Path(__file__).resolve().parents[1]
    out = repo_root / "test_big.xlsx"
    if out.exists():
        out.unlink()
    print(f"Creating large workbook at: {out}")
    make_workbook(out, sheets=5, rows_per_sheet=2000)

    measure_load(out, read_only=False)
    measure_load(out, read_only=True)


if __name__ == "__main__":
    main()
